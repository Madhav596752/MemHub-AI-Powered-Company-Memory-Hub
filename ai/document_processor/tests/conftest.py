"""Pytest fixtures and sample test dataset generation for MemHub."""

import csv
from pathlib import Path
import pytest
import docx
import pypdf

# Directory for sample test files
SAMPLE_DATA_DIR = Path(__file__).parent / "sample_data"


def generate_sample_pdf(file_path: Path, pages_text: list):
    """Generates a valid PDF with text across multiple pages using pypdf."""
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = pypdf.PdfWriter()
    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font)

    for text in pages_text:
        page = writer.add_blank_page(width=612, height=792)
        resources = DictionaryObject()
        fonts = DictionaryObject()
        fonts[NameObject("/F1")] = font_ref
        resources[NameObject("/Font")] = fonts
        page[NameObject("/Resources")] = resources

        stream = DecodedStreamObject()
        stream_data = f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET\n".encode("latin-1", errors="replace")
        stream.set_data(stream_data)
        page[NameObject("/Contents")] = stream

    with open(file_path, "wb") as f:
        writer.write(f)


@pytest.fixture(scope="session")
def sample_dataset(tmp_path_factory) -> dict:
    """Prepares and populates sample documents for testing all formats."""
    data_dir = tmp_path_factory.mktemp("test_docs")

    # 1. Plain Text (TXT)
    txt_path = data_dir / "system_architecture.txt"
    txt_content = (
        "MemHub System Architecture Overview\n\n"
        "The system relies on high-throughput microservices communicating over gRPC and REST APIs.\n"
        "Configuration parameters are stored in /etc/memhub/config.yaml and accessed securely.\n"
        "For documentation, refer to https://internal.memhub.io/docs/v2/arch.\n"
        "Memory limits should not exceed MAX_HEAP_SIZE=4096MB.\n"
    )
    txt_path.write_text(txt_content, encoding="utf-8")

    # 2. Markdown (MD)
    md_path = data_dir / "auth_guide.md"
    md_content = (
        "# Authentication & Authorization Guide\n\n"
        "## Overview\n"
        "MemHub enforces OAuth 2.0 with JWT access tokens for organizational identity.\n\n"
        "## Endpoints\n"
        "- POST /api/auth/token: Exchange credentials for a signed JWT.\n"
        "- GET /api/auth/verify: Verify bearer token claims.\n"
        "- DELETE /api/auth/revoke: Invalidate active refresh tokens.\n\n"
        "Tokens expire after 3600 seconds. Ensure the Authorization: Bearer <TOKEN> header is sent.\n"
    )
    md_path.write_text(md_content, encoding="utf-8")

    # 3. CSV
    csv_path = data_dir / "team_roster.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["EmployeeID", "Name", "Department", "Role", "Email"])
        writer.writerow(["E101", "Alice Vance", "Platform", "Staff Engineer", "alice@acme.co"])
        writer.writerow(["E102", "Bob Miller", "Security", "Infra Sec Lead", "bob@acme.co"])
        writer.writerow(["E103", "Carol Danvers", "Data & AI", "ML Researcher", "carol@acme.co"])

    # 4. DOCX
    docx_path = data_dir / "onboarding_handbook.docx"
    doc = docx.Document()
    doc.add_heading("Acme Engineering Onboarding Handbook", level=1)
    doc.add_paragraph("Welcome to the team! This document details development practices and local setup.")
    doc.add_heading("Development Environment", level=2)
    doc.add_paragraph("Install Docker, Node.js 22, and Python 3.10. Run docker compose up to launch dependencies.")
    
    # Add a table
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Service"
    hdr_cells[1].text = "Port"
    hdr_cells[2].text = "Protocol"
    row1 = table.add_row().cells
    row1[0].text = "API Gateway"
    row1[1].text = "3000"
    row1[2].text = "HTTP/REST"
    row2 = table.add_row().cells
    row2[0].text = "Vector Store"
    row2[1].text = "6333"
    row2[2].text = "gRPC"
    doc.save(docx_path)

    # 5. Multi-page PDF
    pdf_path = data_dir / "annual_report.pdf"
    generate_sample_pdf(
        pdf_path,
        [
            "Annual Report 2026 - Executive Summary and Financial Overview for Acme Corp.",
            "Chapter 1 - Operations and Enterprise Knowledge Growth Across 12 Quarters.",
        ],
    )

    # 6. Scanned / Blank PDF (Image-only simulation)
    scanned_pdf_path = data_dir / "scanned_invoice.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with open(scanned_pdf_path, "wb") as f:
        writer.write(f)

    # 7. Empty file (0 bytes)
    empty_path = data_dir / "empty_doc.txt"
    empty_path.write_bytes(b"")

    # 8. Unsupported file
    unsupported_path = data_dir / "soundtrack.mp3"
    unsupported_path.write_bytes(b"ID3\x03\x00\x00\x00\x00")

    # 9. Corrupted PDF file
    corrupted_path = data_dir / "corrupted_file.pdf"
    corrupted_path.write_bytes(b"%PDF-1.4 incomplete garbage bytes that cannot be parsed \xff\xfe\x00\x01")

    return {
        "txt": txt_path,
        "md": md_path,
        "csv": csv_path,
        "docx": docx_path,
        "pdf": pdf_path,
        "scanned_pdf": scanned_pdf_path,
        "empty": empty_path,
        "unsupported": unsupported_path,
        "corrupted": corrupted_path,
    }
