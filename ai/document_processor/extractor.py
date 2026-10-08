"""Document extraction engine supporting PDF, DOCX, TXT, MD, and CSV.

Extracts structured sections with page numbers, tabular mappings, and error checks.
"""

import csv
import io
import os
from pathlib import Path
from typing import List, Optional, Union

from ai.document_processor.exceptions import (
    UnsupportedFileTypeError,
    EmptyDocumentError,
    ScannedPDFError,
    CorruptedFileError,
    ExtractionError,
)
from ai.document_processor.schemas import ExtractedDocument, ExtractedSection
from ai.document_processor.cleaner import clean_text


class BaseExtractor:
    """Base interface for format-specific extractors."""

    def extract(self, file_path: Path) -> ExtractedDocument:
        raise NotImplementedError


class PDFExtractor(BaseExtractor):
    """Extracts text page-by-page from PDF files using pypdf."""

    MIN_TOTAL_CHARS_THRESHOLD = 30
    MIN_AVG_CHARS_PER_PAGE = 8

    def extract(self, file_path: Path) -> ExtractedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ExtractionError(f"pypdf is required for PDF extraction: {e}")

        sections: List[ExtractedSection] = []
        total_text_length = 0

        try:
            reader = PdfReader(str(file_path))
            total_pages = len(reader.pages)
        except Exception as e:
            raise CorruptedFileError(f"Could not parse or open PDF '{file_path.name}': {e}")

        if total_pages == 0:
            raise EmptyDocumentError(f"PDF file '{file_path.name}' contains 0 pages.")

        for page_idx, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as e:
                page_text = ""
            
            cleaned_page = clean_text(page_text)
            total_text_length += len(cleaned_page)

            if cleaned_page:
                sections.append(
                    ExtractedSection(
                        text=cleaned_page,
                        page=page_idx,
                        metadata={"page_number": page_idx, "page_characters": len(cleaned_page)},
                    )
                )

        # Scanned PDF detection: check whether readable textual content is negligible
        avg_chars = total_text_length / max(total_pages, 1)
        if total_text_length < self.MIN_TOTAL_CHARS_THRESHOLD or avg_chars < self.MIN_AVG_CHARS_PER_PAGE:
            raise ScannedPDFError(
                f"PDF '{file_path.name}' appears to be a scanned document or image-only PDF "
                f"({total_text_length} text characters across {total_pages} pages).",
                details={"total_pages": total_pages, "extracted_characters": total_text_length},
            )

        return ExtractedDocument(
            filename=file_path.name,
            file_type="PDF",
            sections=sections,
            total_pages=total_pages,
            metadata={"total_pages": total_pages, "extracted_pages": len(sections)},
        )


class DOCXExtractor(BaseExtractor):
    """Extracts text, headings, and tables from DOCX files using python-docx."""

    def extract(self, file_path: Path) -> ExtractedDocument:
        try:
            import docx
        except ImportError as e:
            raise ExtractionError(f"python-docx is required for DOCX extraction: {e}")

        sections: List[ExtractedSection] = []

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise CorruptedFileError(f"Could not parse DOCX file '{file_path.name}': {e}")

        accumulated_paragraphs: List[str] = []

        # Extract paragraphs with heading markers
        for p in doc.paragraphs:
            text = clean_text(p.text)
            if not text:
                continue

            style_name = (p.style.name if p.style else "").lower()
            if "heading 1" in style_name:
                text = f"# {text}"
            elif "heading 2" in style_name:
                text = f"## {text}"
            elif "heading 3" in style_name:
                text = f"### {text}"

            accumulated_paragraphs.append(text)

        # Also extract structured content from tables
        for table_idx, table in enumerate(doc.tables, start=1):
            table_rows = []
            header_cells = []
            for row_idx, row in enumerate(table.rows):
                cells = [clean_text(cell.text) for cell in row.cells]
                # deduplicate consecutive identical cells from merged cells
                deduped = []
                for c in cells:
                    if not deduped or deduped[-1] != c:
                        deduped.append(c)

                if row_idx == 0:
                    header_cells = deduped
                    table_rows.append(f"Table {table_idx} Header: " + " | ".join(header_cells))
                else:
                    if header_cells and len(header_cells) == len(deduped):
                        row_repr = " | ".join(f"[{h}]: {val}" for h, val in zip(header_cells, deduped) if val)
                    else:
                        row_repr = " | ".join(val for val in deduped if val)
                    if row_repr:
                        table_rows.append(f"Row {row_idx}: {row_repr}")

            if table_rows:
                accumulated_paragraphs.append("\n".join(table_rows))

        full_content = "\n\n".join(accumulated_paragraphs)
        if not full_content.strip():
            raise EmptyDocumentError(f"DOCX file '{file_path.name}' is empty.")

        sections.append(
            ExtractedSection(
                text=full_content,
                page=None,
                metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
            )
        )

        return ExtractedDocument(
            filename=file_path.name,
            file_type="DOCX",
            sections=sections,
            total_pages=None,
            metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
        )


class TextExtractor(BaseExtractor):
    """Extracts plain text and Markdown files with robust encoding fallback."""

    ENCODINGS = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]

    def __init__(self, file_type: str = "TXT"):
        self.file_type = file_type

    def extract(self, file_path: Path) -> ExtractedDocument:
        content = None
        used_encoding = None

        for enc in self.ENCODINGS:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                    used_encoding = enc
                    break
            except (UnicodeDecodeError, UnicodeError):
                continue
            except Exception as e:
                raise CorruptedFileError(f"Failed to read file '{file_path.name}': {e}")

        if content is None:
            # Fallback with replacement
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    content = f.read()
                    used_encoding = "utf-8-replace"
            except Exception as e:
                raise CorruptedFileError(f"Failed to read file '{file_path.name}': {e}")

        cleaned = clean_text(content)
        if not cleaned:
            raise EmptyDocumentError(f"File '{file_path.name}' is empty.")

        sections = [
            ExtractedSection(
                text=cleaned,
                page=None,
                metadata={"encoding": used_encoding, "character_count": len(cleaned)},
            )
        ]

        return ExtractedDocument(
            filename=file_path.name,
            file_type=self.file_type,
            sections=sections,
            total_pages=None,
            metadata={"encoding": used_encoding},
        )


class CSVExtractor(BaseExtractor):
    """Extracts CSV data converting rows into semantically meaningful text records.
    
    Preserves column names and structure so embeddings retain key-value context.
    """

    def extract(self, file_path: Path) -> ExtractedDocument:
        text_content = None
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    text_content = f.read()
                    break
            except UnicodeDecodeError:
                continue
            except Exception as e:
                raise CorruptedFileError(f"Failed to read CSV '{file_path.name}': {e}")

        if text_content is None or not text_content.strip():
            raise EmptyDocumentError(f"CSV file '{file_path.name}' is empty.")

        # Detect delimiter and dialect
        try:
            sample = text_content[:4096]
            sniffer = csv.Sniffer()
            delimiter = ","
            try:
                dialect = sniffer.sniff(sample, delimiters=",;\t|")
                delimiter = dialect.delimiter
            except Exception:
                delimiter = ","
        except Exception:
            delimiter = ","

        reader = csv.reader(io.StringIO(text_content), delimiter=delimiter)
        rows = [row for row in reader if any(cell.strip() for cell in row)]

        if not rows:
            raise EmptyDocumentError(f"CSV file '{file_path.name}' has no readable data rows.")

        headers = [clean_text(h) for h in rows[0]]
        data_rows = rows[1:]

        if not data_rows:
            # Header only
            header_line = "CSV Header columns: " + ", ".join(h for h in headers if h)
            sections = [
                ExtractedSection(
                    text=header_line,
                    page=None,
                    metadata={"headers": headers, "row_count": 0},
                )
            ]
            return ExtractedDocument(
                filename=file_path.name,
                file_type="CSV",
                sections=sections,
                metadata={"headers": headers, "total_rows": 0},
            )

        # Convert each row into semantic representation
        semantic_records: List[str] = []
        for row_idx, row in enumerate(data_rows, start=1):
            field_pairs = []
            for col_idx, cell in enumerate(row):
                val = clean_text(cell)
                if not val:
                    continue
                col_name = headers[col_idx] if col_idx < len(headers) and headers[col_idx] else f"Column_{col_idx+1}"
                field_pairs.append(f"[{col_name}]: {val}")

            if field_pairs:
                record_text = f"Record {row_idx}: " + " | ".join(field_pairs)
                semantic_records.append(record_text)

        full_csv_text = "\n".join(semantic_records)
        sections = [
            ExtractedSection(
                text=full_csv_text,
                page=None,
                metadata={"headers": headers, "total_rows": len(data_rows)},
            )
        ]

        return ExtractedDocument(
            filename=file_path.name,
            file_type="CSV",
            sections=sections,
            total_pages=None,
            metadata={"headers": headers, "total_rows": len(data_rows)},
        )


class DocumentExtractor:
    """Unified entry point for extracting any supported document type."""

    SUPPORTED_EXTENSIONS = {
        ".pdf": ("PDF", PDFExtractor),
        ".docx": ("DOCX", DOCXExtractor),
        ".txt": ("TXT", lambda: TextExtractor("TXT")),
        ".md": ("MD", lambda: TextExtractor("MD")),
        ".markdown": ("MD", lambda: TextExtractor("MD")),
        ".csv": ("CSV", CSVExtractor),
    }

    def __init__(self):
        self._extractors = {
            "PDF": PDFExtractor(),
            "DOCX": DOCXExtractor(),
            "TXT": TextExtractor("TXT"),
            "MD": TextExtractor("MD"),
            "CSV": CSVExtractor(),
        }

    def get_file_type(self, file_path: Union[str, Path]) -> str:
        """Determines uppercase file type from extension."""
        path = Path(file_path)
        ext = path.suffix.lower()
        if ext in self.SUPPORTED_EXTENSIONS:
            return self.SUPPORTED_EXTENSIONS[ext][0]
        raise UnsupportedFileTypeError(
            f"Unsupported file format '{ext}' for file '{path.name}'. "
            f"Supported formats: PDF, DOCX, TXT, MD, CSV."
        )

    def extract(self, file_path: Union[str, Path]) -> ExtractedDocument:
        """Extracts text content and metadata from the given file."""
        path = Path(file_path)
        if not path.exists():
            raise CorruptedFileError(f"File not found: '{path}'")

        if path.is_file() and path.stat().st_size == 0:
            raise EmptyDocumentError(f"File '{path.name}' is 0 bytes (empty).")

        file_type = self.get_file_type(path)
        extractor = self._extractors[file_type]
        return extractor.extract(path)
